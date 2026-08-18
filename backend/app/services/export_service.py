"""Word 文档导出服务."""

import io

from app.services.format_spec import FormatSpec, build_format_spec
from app.services.storage_service import upload_file


def _apply_margins(doc, spec: FormatSpec) -> None:
    """应用页边距（未指定的边沿用 docx 默认）."""
    if not spec.margins_cm:
        return
    from docx.shared import Cm

    section = doc.sections[0]
    attr_map = {
        "top": "top_margin",
        "bottom": "bottom_margin",
        "left": "left_margin",
        "right": "right_margin",
    }
    for side, attr in attr_map.items():
        if side in spec.margins_cm:
            setattr(section, attr, Cm(spec.margins_cm[side]))


def _apply_body_format(para, spec: FormatSpec) -> None:
    """应用正文行距、字体、字号（含中文 eastAsia 字体）."""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    if spec.line_spacing_fixed_pt:
        para.paragraph_format.line_spacing = Pt(spec.line_spacing_fixed_pt)
    elif spec.line_spacing:
        para.paragraph_format.line_spacing = spec.line_spacing
    for run in para.runs:
        run.font.size = Pt(spec.body_size_pt)
        if spec.body_font:
            run.font.name = spec.body_font
            run._element.rPr.rFonts.set(qn("w:eastAsia"), spec.body_font)


def _apply_heading_size(heading, size_pt: float | None) -> None:
    """一级章节标题按格式要求调字号，其余沿用 heading 样式."""
    if not size_pt:
        return
    from docx.shared import Pt

    for run in heading.runs:
        run.font.size = Pt(size_pt)


async def export_to_word(
    chapters: dict[str, str],
    outline: list[dict],
    project_name: str = "技术方案",
    format_requirements: list[dict] | None = None,
) -> str:
    """将章节内容导出为 Word 文档，返回 MinIO storage_key.

    format_requirements 来自招标文件 meta，经 format_spec 解析后驱动排版；
    未提供或不可解析时回退默认样式（正文 12pt、行距 1.5 倍、默认边距）。
    """
    try:
        from docx import Document
    except ImportError:
        from app.core.exceptions import BizError

        raise BizError(code=5010, message="python-docx 未安装") from None

    spec = build_format_spec(format_requirements)
    doc = Document()
    _apply_margins(doc, spec)

    # 标题
    title = doc.add_heading(project_name, level=0)
    title.alignment = 1  # 居中

    # 按大纲顺序写入章节
    for chapter_info in outline:
        chapter_no = chapter_info.get("chapter_no", "")
        chapter_title = chapter_info.get("title", "")
        content = chapters.get(chapter_no, "")

        # 章节标题
        heading = doc.add_heading(f"{chapter_no} {chapter_title}", level=1)
        _apply_heading_size(heading, spec.heading_size_pt)

        # 章节内容（按段落拆分）
        for paragraph_text in content.split("\n\n"):
            if paragraph_text.strip():
                # 检查是否是子标题（以 ## 开头）
                stripped = paragraph_text.strip()
                if stripped.startswith("## "):
                    doc.add_heading(stripped[3:], level=2)
                elif stripped.startswith("### "):
                    doc.add_heading(stripped[4:], level=3)
                else:
                    para = doc.add_paragraph(stripped)
                    _apply_body_format(para, spec)

    # 保存到内存
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    # 上传到 MinIO
    filename = f"{project_name}.docx"
    storage_key = upload_file(
        file_content=buffer,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    return storage_key
