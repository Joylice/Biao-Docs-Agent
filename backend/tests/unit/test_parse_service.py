"""招标解析服务测试."""

import io

import pytest

from app.core.exceptions import BizError
from app.services.parse_service import (
    _scoring_anchor,
    extract_tender_text,
    select_parse_window,
)


def _make_pdf_bytes(text: str = "Tender Document Sample Text") -> bytes:
    """构造最小合法 PDF（单页 + Helvetica 文本流，xref 偏移程序内计算）."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    """用 python-docx 构造最小 DOCX bytes."""
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_docx_with_table(paragraphs: list[str], table_rows: list[list[str]]) -> bytes:
    """构造含表格的 DOCX bytes."""
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
    for i, row in enumerate(table_rows):
        for j, cell_text in enumerate(row):
            table.cell(i, j).text = cell_text
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestExtractTenderText:
    """文本提取测试."""

    @pytest.mark.asyncio
    async def test_txt_extraction(self) -> None:
        """TXT 文件直接解码."""
        content = "测试文本内容".encode()
        result = await extract_tender_text(content, "test.txt")
        assert "测试文本内容" in result

    @pytest.mark.asyncio
    async def test_unsupported_format_raises(self) -> None:
        """不支持的格式抛出 BizError."""
        with pytest.raises(BizError, match=r"不支持的文件格式"):
            await extract_tender_text(b"fake content", "test.xyz")

    @pytest.mark.asyncio
    async def test_pdf_without_library(self) -> None:
        """PDF 解析在无库时抛出明确错误."""
        try:
            import pdfplumber  # noqa: F401

            pytest.skip("pdfplumber 已安装，无法测试缺失场景")
        except ImportError:
            with pytest.raises(BizError, match=r"pdfplumber 未安装|PDF 解析失败"):
                await extract_tender_text(b"fake pdf", "test.pdf")

    @pytest.mark.asyncio
    async def test_docx_without_library(self) -> None:
        """Word 解析在无库时抛出明确错误."""
        try:
            import docx  # noqa: F401

            pytest.skip("python-docx 已安装，无法测试缺失场景")
        except ImportError:
            with pytest.raises(BizError, match=r"python-docx 未安装|Word 解析失败"):
                await extract_tender_text(b"fake docx", "test.docx")

    @pytest.mark.asyncio
    async def test_pdf_extraction_with_real_bytes(self) -> None:
        """真实最小 PDF bytes 能解析出非空文本（回归：pdfplumber.open 位置传参）."""
        content = _make_pdf_bytes("Tender Document Sample Text")
        result = await extract_tender_text(content, "招标文件.pdf")
        assert result.strip(), "PDF 解析结果不应为空"
        assert "Tender Document Sample Text" in result

    @pytest.mark.asyncio
    async def test_docx_extraction_with_real_bytes(self) -> None:
        """真实最小 DOCX bytes 能解析出非空文本（回归：docx.Document 位置传参）."""
        content = _make_docx_bytes(["招标公告正文第一段", "技术要求清单第二段"])
        result = await extract_tender_text(content, "招标文件.docx")
        assert result.strip(), "DOCX 解析结果不应为空"
        assert "招标公告正文第一段" in result
        assert "技术要求清单第二段" in result

    @pytest.mark.asyncio
    async def test_docx_table_content_extracted(self) -> None:
        """DOCX 表格内容必须被提取：评分标准/技术需求常置于表格，只取段落会丢失."""
        content = _make_docx_with_table(
            paragraphs=["第一章 招标公告"],
            table_rows=[
                ["评分项", "分值"],
                ["技术方案完整性", "60"],
                ["报价", "40"],
            ],
        )
        result = await extract_tender_text(content, "招标文件.docx")
        assert "技术方案完整性" in result, "表格内的评分项文本丢失"
        assert "报价" in result


class TestSelectParseWindow:
    """长文本窗口选择测试（截断必须覆盖评分/技术需求区，而非固定取前 N 字符）."""

    def test_short_text_unchanged(self) -> None:
        """不超预算的文本原样返回."""
        text = "短文本：评分标准 技术方案 60 分"
        assert select_parse_window(text, budget=1000) == text

    def test_window_covers_scoring_section(self) -> None:
        """超长文本的窗口必须包含评分标准区（原实现取前 8000 字符导致评分点在窗口外）."""
        filler = "目录内容占位。" * 2000  # 前置大量目录/封面内容
        scoring = "评分标准：技术方案 60 分，报价 40 分。" * 50
        tail = "合同条款占位。" * 3000
        text = filler + scoring + tail
        window = select_parse_window(text, budget=2000)
        assert len(window) <= 2000
        assert "评分标准" in window, "窗口未覆盖评分标准区"

    def test_window_keeps_head_for_project_name(self) -> None:
        """窗口保留文档开头（项目名称/编号位于封面）."""
        head = "项目名称：智慧平台建设项目 编号 TN-001。"
        filler = "目录内容占位。" * 2000
        scoring = "评分标准：技术方案 60 分。" * 50
        text = head + filler + scoring
        window = select_parse_window(text, budget=3000)
        assert "项目名称" in window
        assert "评分标准" in window

    def test_no_scoring_keyword_falls_back_to_head(self) -> None:
        """无评分关键词时回退取头部预算字符."""
        text = "普通文本占位。" * 3000
        window = select_parse_window(text, budget=1000)
        assert len(window) <= 1000
        assert window == text[:1000]

    def test_window_covers_both_scoring_and_tech_regions(self) -> None:
        """评分细则与技术要求分处文档不同区域时，窗口必须同时覆盖两处.

        回归：真实招标文件 19.7 万字符，单窗口锚定前附表时细则/技要求落窗外，
        LLM 只能给出“详见招标文件”类笼统 criteria。
        """
        filler_head = "封面目录占位。" * 1000
        scoring = "评分标准：施工方案合理性 30 分，工期安排 20 分。" * 80
        filler_mid = "合同条款占位。" * 3000
        tech = "技术要求：系统可用率 99.9%，支持双机热备。" * 80
        filler_tail = "附件占位。" * 1000
        text = filler_head + scoring + filler_mid + tech + filler_tail
        window = select_parse_window(text, budget=4000)
        assert len(window) <= 4000
        assert "施工方案合理性" in window, "窗口未覆盖评分细则区"
        assert "双机热备" in window, "窗口未覆盖技术要求区"

    def test_anchor_prefers_scoring_criteria_over_front_table(self) -> None:
        """锚点优先取评分标准/评分因素细则区，而非评标办法前附表（前附表只有汇总无细则）.

        回归：真实文档中评标办法前附表（12656）早于评分因素细则（29084），
        锚定前附表时窗口被前附表+投标人须知占满，细则落窗外。
        """
        head = "封面占位。" * 200
        front_table = "评标办法前附表：条款号 内容。" * 500
        filler = "投标人须知占位。" * 2000
        criteria = "评分标准：施工方案 30 分，报价 50 分。" * 80
        text = head + front_table + filler + criteria
        anchor = _scoring_anchor(text)
        assert anchor == text.find("评分标准"), "应锚定评分标准细则区而非评标办法前附表"

    def test_anchor_falls_back_to_generic_scoring_words(self) -> None:
        """无细则关键词时回退到评标办法/评分等通用词."""
        text = "封面。" * 100 + "评标办法前附表：技术 60 分，报价 40 分。"
        anchor = _scoring_anchor(text)
        assert anchor == text.find("评标办法")


class TestParseTenderWithLlmSchema:
    """重新解析（include_tech_requirements=False）时输出 schema 不含技术需求."""

    @pytest.mark.asyncio
    async def test_schema_excludes_tech_requirements_when_disabled(self, monkeypatch) -> None:
        from app.services import llm_service, parse_service

        captured: dict = {}

        async def fake_call(
            system_prompt: str,
            user_prompt: str,
            response_format: dict | None = None,
            mock: bool | None = None,
        ) -> dict:
            captured["schema"] = (response_format or {}).get("json_schema", {}).get("schema", {})
            # 真实 LLM 可能无视 schema 裁剪照常输出技术需求（提示词正文仍含指令）
            return {
                "score_points": [{"clause_no": "1", "item": "方案"}],
                "tech_requirements": [{"seq": 1, "description": "不应落库"}],
            }

        monkeypatch.setattr(llm_service, "call_llm_with_schema", fake_call)

        parsed = await parse_service.parse_tender_with_llm(
            "招标正文", include_tech_requirements=False
        )
        schema = captured["schema"]
        assert "tech_requirements" not in schema.get("properties", {})
        assert "tech_requirements" not in schema.get("required", [])
        assert "score_points" in schema.get("properties", {})
        assert parsed.tech_requirements == [], "禁用时 LLM 输出的技术需求必须丢弃"

    @pytest.mark.asyncio
    async def test_schema_includes_tech_requirements_by_default(self, monkeypatch) -> None:
        from app.services import llm_service, parse_service

        captured: dict = {}

        async def fake_call(
            system_prompt: str,
            user_prompt: str,
            response_format: dict | None = None,
            mock: bool | None = None,
        ) -> dict:
            captured["schema"] = (response_format or {}).get("json_schema", {}).get("schema", {})
            return {
                "score_points": [{"clause_no": "1", "item": "方案"}],
                "tech_requirements": [{"seq": 1, "description": "高可用"}],
            }

        monkeypatch.setattr(llm_service, "call_llm_with_schema", fake_call)

        parsed = await parse_service.parse_tender_with_llm("招标正文")
        schema = captured["schema"]
        assert "tech_requirements" in schema.get("properties", {})
        assert parsed.tech_requirements == [{"seq": 1, "description": "高可用"}]


class TestFormatRequirementsSchema:
    """格式要求提取：首次/重新解析均提取，缺失时兜底空列表."""

    @pytest.mark.asyncio
    async def test_schema_includes_format_requirements(self, monkeypatch) -> None:
        from app.services import llm_service, parse_service

        captured: dict = {}

        async def fake_call(
            system_prompt: str,
            user_prompt: str,
            response_format: dict | None = None,
            mock: bool | None = None,
        ) -> dict:
            captured["schema"] = (response_format or {}).get("json_schema", {}).get("schema", {})
            return {
                "score_points": [{"clause_no": "1", "item": "方案"}],
                "format_requirements": [{"category": "font_body", "requirement": "正文宋体小四"}],
            }

        monkeypatch.setattr(llm_service, "call_llm_with_schema", fake_call)

        parsed = await parse_service.parse_tender_with_llm("招标正文")
        props = captured["schema"]["properties"]
        assert "format_requirements" in props
        item_schema = props["format_requirements"]["items"]
        assert set(item_schema["required"]) == {"category", "requirement"}
        assert parsed.format_requirements == [
            {"category": "font_body", "requirement": "正文宋体小四"}
        ]

    @pytest.mark.asyncio
    async def test_format_requirements_kept_when_score_points_only(self, monkeypatch) -> None:
        """重新解析（跳过技术需求）仍提取格式要求."""
        from app.services import llm_service, parse_service

        async def fake_call(
            system_prompt: str,
            user_prompt: str,
            response_format: dict | None = None,
            mock: bool | None = None,
        ) -> dict:
            return {
                "score_points": [{"clause_no": "1", "item": "方案"}],
                "format_requirements": [{"category": "margin", "requirement": "左边距3cm"}],
            }

        monkeypatch.setattr(llm_service, "call_llm_with_schema", fake_call)

        parsed = await parse_service.parse_tender_with_llm(
            "招标正文", include_tech_requirements=False
        )
        assert parsed.tech_requirements == []
        assert parsed.format_requirements == [{"category": "margin", "requirement": "左边距3cm"}]

    @pytest.mark.asyncio
    async def test_missing_format_requirements_fallback_empty(self, monkeypatch) -> None:
        from app.services import llm_service, parse_service

        async def fake_call(
            system_prompt: str,
            user_prompt: str,
            response_format: dict | None = None,
            mock: bool | None = None,
        ) -> dict:
            return {"score_points": [{"clause_no": "1", "item": "方案"}]}

        monkeypatch.setattr(llm_service, "call_llm_with_schema", fake_call)

        parsed = await parse_service.parse_tender_with_llm("招标正文")
        assert parsed.format_requirements == []


class TestSaveParseResultFormatRequirements:
    """格式要求落库：写入招标文件 Document.meta.format_requirements."""

    @pytest.mark.asyncio
    async def test_save_writes_format_requirements_to_meta(self) -> None:
        import uuid
        from unittest.mock import AsyncMock, MagicMock

        from app.models.document import Document
        from app.services.parse_service import ParsedTender, save_parse_result

        doc_id = uuid.uuid4()
        project_id = uuid.uuid4()
        doc = Document(
            id=doc_id,
            project_id=project_id,
            doc_type="tender_file",
            title="t.docx",
            storage_key="k",
            status="parsing",
        )
        doc.meta = {}

        db = MagicMock()
        db.flush = AsyncMock()
        result = MagicMock()
        result.scalar_one_or_none.return_value = doc
        db.execute = AsyncMock(return_value=result)

        fr = [{"category": "line_spacing", "requirement": "1.5倍行距"}]
        parsed = ParsedTender(
            score_points=[{"clause_no": "1", "item": "方案"}],
            tech_requirements=[],
            format_requirements=fr,
        )
        sp_count, tr_count = await save_parse_result(db, project_id, doc_id, parsed)

        assert (sp_count, tr_count) == (1, 0)
        assert doc.meta["format_requirements"] == fr
        assert doc.status == "parsed"
