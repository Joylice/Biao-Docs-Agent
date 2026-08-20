"""export_service 单元测试 — 格式要求驱动 Word 排版.

覆盖：默认样式（正文 12pt）、格式要求应用（字体/字号/行距/页边距/标题字号），
通过捕获上传 buffer 重新打开 docx 断言。
"""

import base64
import io

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

import app.services.export_service as export_service
from app.core.exceptions import BizError

OUTLINE = [{"chapter_no": "1", "title": "项目概述"}]
CHAPTERS = {"1": "## 背景\n\n这是正文段落内容，用于断言排版属性。"}

# 1x1 合法 PNG（内嵌断言用）
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


@pytest.fixture
def capture_upload(monkeypatch):
    """拦截 MinIO 上传，返回捕获的 buffer 容器."""
    captured: dict = {}

    def fake_upload(file_content, filename, content_type):
        captured["buffer"] = file_content
        captured["filename"] = filename
        return "test/export.docx"

    monkeypatch.setattr(export_service, "upload_file", fake_upload)
    return captured


def _open(captured: dict) -> Document:
    buffer: io.BytesIO = captured["buffer"]
    buffer.seek(0)
    return Document(buffer)


def _body_paragraph(doc: Document):
    """取正文段落（Normal 样式，排除标题/Title/封面/目录域空段）."""
    return next(
        p
        for p in doc.paragraphs
        if p.style.name == "Normal" and p.text.strip() and "编制日期" not in p.text
    )


class TestDefaultStyle:
    @pytest.mark.asyncio
    async def test_export_defaults_without_requirements(self, capture_upload) -> None:
        """无格式要求：正文 12pt、行距 1.5 倍、默认页边距."""
        await export_service.export_to_word(CHAPTERS, OUTLINE, "测试项目")
        doc = _open(capture_upload)
        para = _body_paragraph(doc)
        assert para.runs[0].font.size == Pt(12)
        assert para.paragraph_format.line_spacing == 1.5


class TestFormatApplied:
    @pytest.mark.asyncio
    async def test_export_applies_format_requirements(self, capture_upload) -> None:
        """格式要求应用：宋体四号正文、三号章标题、固定行距、页边距."""
        fr = [
            {"category": "font_body", "requirement": "正文宋体四号"},
            {"category": "font_heading", "requirement": "章标题三号"},
            {"category": "line_spacing", "requirement": "固定值28磅"},
            {"category": "margin", "requirement": "上下2.54cm，左右3.17cm"},
        ]
        await export_service.export_to_word(CHAPTERS, OUTLINE, "测试项目", fr)
        doc = _open(capture_upload)

        # 正文字体/字号（含 eastAsia 中文字体）
        para = _body_paragraph(doc)
        run = para.runs[0]
        assert run.font.size == Pt(14)
        assert run.font.name == "宋体"
        assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"

        # 固定行距
        assert para.paragraph_format.line_spacing == Pt(28)

        # 一级章节标题字号
        h1 = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
        assert h1.runs[0].font.size == Pt(16)

        # 页边距（docx 以 twip 存储，按 twip 粒度断言）
        section = doc.sections[0]
        assert section.top_margin == Twips(1440)
        assert section.bottom_margin == Twips(1440)
        assert section.left_margin == Twips(1797)
        assert section.right_margin == Twips(1797)

    @pytest.mark.asyncio
    async def test_export_unparsable_requirements_fallback(self, capture_upload) -> None:
        """不可解析的格式要求回退默认样式，不阻塞导出."""
        fr = [{"category": "font_body", "requirement": "排版美观"}]
        key = await export_service.export_to_word(CHAPTERS, OUTLINE, "测试项目", fr)
        assert key == "test/export.docx"
        doc = _open(capture_upload)
        assert _body_paragraph(doc).runs[0].font.size == Pt(12)


class TestStageEEnhancements:
    """阶段 E4：封面/目录域/页码页脚/分页控制/对标附表/引用标注."""

    @pytest.mark.asyncio
    async def test_cover_page(self, capture_upload) -> None:
        """封面：项目名标题 + 公司名 + 编制日期占位."""
        await export_service.export_to_word(
            CHAPTERS, OUTLINE, "测试项目", company_name="投标公司"
        )
        doc = _open(capture_upload)
        texts = [p.text for p in doc.paragraphs]
        assert texts[0] == "测试项目"
        assert any(t == "投标公司" for t in texts)
        assert any(t.startswith("编制日期：") for t in texts)

    @pytest.mark.asyncio
    async def test_toc_field(self, capture_upload) -> None:
        """目录域：TOC field 写入（打开时更新）."""
        await export_service.export_to_word(CHAPTERS, OUTLINE, "测试项目")
        doc = _open(capture_upload)
        xml = doc.element.xml
        assert "TOC \\o" in xml
        assert any(p.text == "目录" for p in doc.paragraphs)

    @pytest.mark.asyncio
    async def test_footer_page_number_field(self, capture_upload) -> None:
        """页脚页码域：PAGE field."""
        await export_service.export_to_word(CHAPTERS, OUTLINE, "测试项目")
        doc = _open(capture_upload)
        footer_xml = doc.sections[0].footer._element.xml
        assert "PAGE" in footer_xml
        assert "fldChar" in footer_xml

    @pytest.mark.asyncio
    async def test_chapter_page_break_before(self, capture_upload) -> None:
        """分页控制：章节标题前分页（每章起始新页）."""
        await export_service.export_to_word(CHAPTERS, OUTLINE, "测试项目")
        doc = _open(capture_upload)
        h1 = next(
            p
            for p in doc.paragraphs
            if p.style.name == "Heading 1" and "项目概述" in p.text
        )
        assert h1.paragraph_format.page_break_before is True

    @pytest.mark.asyncio
    async def test_benchmark_appendix_table(self, capture_upload) -> None:
        """附表：评分对标表（阶段 D benchmark 数据）."""
        rows = [
            {
                "clause_no": "2.1",
                "item": "性能",
                "score": 8.0,
                "coverage": 0.3,
                "risk": "high",
                "strategy": "突出等保三级合规设计",
            }
        ]
        await export_service.export_to_word(
            CHAPTERS, OUTLINE, "测试项目", benchmark_rows=rows
        )
        doc = _open(capture_upload)
        assert len(doc.tables) == 1
        table = doc.tables[0]
        header = [c.text for c in table.rows[0].cells]
        assert "条款号" in header and "应对策略" in header
        cells = [c.text for c in table.rows[1].cells]
        assert "2.1" in cells and "突出等保三级合规设计" in cells

    @pytest.mark.asyncio
    async def test_citations_inline(self, capture_upload) -> None:
        """引用标注内联：【来源：{doc_title} P{page_no}】."""
        cites = {"1": [{"chunk_id": "c1", "doc_title": "公司资质材料", "page_no": 3}]}
        await export_service.export_to_word(
            CHAPTERS, OUTLINE, "测试项目", citations_by_chapter=cites
        )
        doc = _open(capture_upload)
        texts = [p.text for p in doc.paragraphs]
        assert any("【来源：公司资质材料 P3】" in t for t in texts)

    @pytest.mark.asyncio
    async def test_citation_without_page(self, capture_upload) -> None:
        """page_no 缺失时省略页码."""
        cites = {"1": [{"chunk_id": "c2", "doc_title": "历史方案", "page_no": None}]}
        await export_service.export_to_word(
            CHAPTERS, OUTLINE, "测试项目", citations_by_chapter=cites
        )
        doc = _open(capture_upload)
        texts = [p.text for p in doc.paragraphs]
        assert any("【来源：历史方案】" in t for t in texts)


class TestImageEmbed:
    """Markdown 图片语法内嵌 Word（阶段 3）."""

    @pytest.mark.asyncio
    async def test_image_embedded_from_signed_url(self, capture_upload, monkeypatch) -> None:
        """签名 URL → 提取 storage_key 拉取字节 → inline shape 内嵌."""
        downloaded: list[str] = []

        def fake_download(key: str) -> bytes:
            downloaded.append(key)
            return PNG_1PX

        monkeypatch.setattr(export_service, "download_file", fake_download)
        url = "http://minio:9000/bid-documents/images/pid/u1/arch.png?X-Amz-Signature=abc"
        chapters = {"1": f"架构示意如下\n\n![架构图]({url})\n\n结尾段落"}
        await export_service.export_to_word(chapters, OUTLINE, "测试项目")
        assert downloaded == ["images/pid/u1/arch.png"]
        doc = _open(capture_upload)
        assert len(doc.inline_shapes) == 1

    @pytest.mark.asyncio
    async def test_image_fetch_failure_fallback_text(self, capture_upload, monkeypatch) -> None:
        """拉取失败降级为文本说明，不阻塞导出."""

        def fake_download(key: str) -> bytes:
            raise BizError(code=5003, message="下载失败")

        monkeypatch.setattr(export_service, "download_file", fake_download)
        chapters = {"1": "![架构图](images/pid/u1/arch.png)"}
        key = await export_service.export_to_word(chapters, OUTLINE, "测试项目")
        assert key == "test/export.docx"
        doc = _open(capture_upload)
        assert len(doc.inline_shapes) == 0
        texts = [p.text for p in doc.paragraphs]
        assert any("[图片: 架构图]" in t for t in texts)

